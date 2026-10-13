"""v2.3: write 1_Google_Sheet/App_Data.json - the release's data tabs as plain text, read by the sheet's
"Proof Generator > Update from GitHub" menu (Updater.gs).

  tabs      Component_Registry, Component_Schema, Product_Assembly, Detail_Drawings, Design_Guide:
            every row as text, exactly as the sheet shows it (these tabs are replaced in place).
  template  Proof_Template rows + every older default of each key, so the update can tell a value you
            never edited (offered for update) from your own wording (kept).
  code      the Apps Script files, with the release each one last changed in, for the
            "paste these files" list when the code cannot update itself.

usage: python3 export_app_data.py <version> <package_dir> [<older package dirs> ...]
"""
import hashlib, json, os, re, sys
from openpyxl import load_workbook

MANAGED = ["Component_Registry", "Component_Schema", "Product_Assembly", "Detail_Drawings", "Design_Guide"]
CODE = [("Code", "SERVER_JS", "Code.gs"), ("ProofLibrary", "SERVER_JS", "ProofLibrary.gs"), ("Updater", "SERVER_JS", "Updater.gs"),
        ("Index", "HTML", "Index.html"), ("JavaScript", "HTML", "JavaScript.html"), ("Stylesheet", "HTML", "Stylesheet.html"),
        ("Assets", "HTML", "Assets.html"), ("appsscript", "JSON", "appsscript.json")]
WB = "1_Google_Sheet/Sign_Proof_Specifications_App_Data.xlsx"


def ver(v): return [int(x) for x in v.split(".")]


def text(v):
    if v is None: return ""
    if isinstance(v, bool): return "TRUE" if v else "FALSE"
    if isinstance(v, float) and v.is_integer(): return str(int(v))
    return str(v)


def rows(ws):
    out = [[text(v) for v in r] for r in ws.iter_rows(values_only=True)]
    while out and not any(out[-1]): out.pop()
    w = len(out[0])                                      # header width
    return [r[:w] + [""] * (w - len(r)) for r in out]


def digest(path):
    return hashlib.sha1(open(path, "rb").read()).hexdigest() if os.path.exists(path) else None


def export(V, pkg, older):
    wb = load_workbook(os.path.join(pkg, WB))
    tabs = {t: rows(wb[t]) for t in MANAGED}
    tpl = rows(wb["Proof_Template"])
    cur = {r[0]: r[1] if len(r) > 1 else "" for r in tpl[1:] if r and r[0]}
    previous = {}
    older = sorted(older, key=lambda d: ver(re.search(r"v([\d.]+)$", d.rstrip("/")).group(1)), reverse=True)
    for d in older:
        p = os.path.join(d, WB)
        if not os.path.exists(p): continue
        old = load_workbook(p, read_only=True)
        if "Proof_Template" not in old.sheetnames: continue
        for r in old["Proof_Template"].iter_rows(min_row=2, values_only=True):
            if not r or not r[0] or r[0] not in cur: continue
            v = text(r[1] if len(r) > 1 else "")
            if v != cur[r[0]] and v not in previous.setdefault(r[0], []): previous[r[0]].append(v)
    previous = {k: v for k, v in previous.items() if v}

    files = []
    for name, typ, fn in CODE:
        path = os.path.join(pkg, "2_Apps_Script", fn)
        assert os.path.exists(path), path
        h, since = digest(path), V
        for d in older:                                  # newest first: walk back while unchanged
            if digest(os.path.join(d, "2_Apps_Script", fn)) != h: break
            since = re.search(r"v([\d.]+)$", d.rstrip("/")).group(1)
        files.append({"name": name, "type": typ, "path": "2_Apps_Script/" + fn, "since": since})

    data = {"version": V, "tabs": tabs, "template": {"rows": tpl, "previous": previous}, "code": {"files": files}}
    out = os.path.join(pkg, "1_Google_Sheet", "App_Data.json")
    with open(out, "w") as f: json.dump(data, f, ensure_ascii=False, separators=(",", ":"))
    return out, data


if __name__ == "__main__":
    out, data = export(sys.argv[1], sys.argv[2], sys.argv[3:])
    print(out, os.path.getsize(out), "bytes;", {t: len(r) - 1 for t, r in data["tabs"].items()}, data["template"]["previous"].keys(),
          [(f["name"], f["since"]) for f in data["code"]["files"]])
